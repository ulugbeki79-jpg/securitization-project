"""Write the waterfall results to a multi-sheet Excel workbook."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill, numbers
from openpyxl.utils import get_column_letter


# UZS number format with thousands separator, no decimals
_UZS_FMT = '#,##0'
_UZS_FMT2 = '#,##0.00'
_PCT_FMT = '0.00%'

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_BOLD = Font(bold=True, size=10)


def _style_sheet(ws, df: pd.DataFrame, uzs_cols: set[str], pct_cols: set[str] | None = None):
    """Apply formatting to a worksheet written from a DataFrame."""
    pct_cols = pct_cols or set()

    # Header row
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Data formatting
    for col_idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        if col_name in uzs_cols:
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = _UZS_FMT
        elif col_name in pct_cols:
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = _PCT_FMT

    # Auto-width (approximate)
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(len(str(col_name)), 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    ws.freeze_panes = "B2"


_POOL_UZS = {
    "beginning_balance", "scheduled_pmt", "scheduled_principal",
    "scheduled_interest", "prepayment", "default_amount", "loss",
    "recovery", "servicing_fee", "net_interest", "total_principal",
    "ending_balance",
}

_TRANCHE_UZS = {
    "beginning_balance", "interest_due", "interest_paid",
    "interest_shortfall", "principal_paid", "loss_allocated",
    "ending_balance", "total_cash_flow", "excess_spread",
}


def write_waterfall_excel(
    pool_cf: pd.DataFrame,
    class_a_cf: pd.DataFrame,
    class_b_cf: pd.DataFrame,
    summary: list[dict],
    assumptions_text: list[str],
    out_path: Path,
) -> Path:
    """Write a formatted multi-sheet Excel workbook.

    Sheets
    ------
    1. Summary           — deal parameters and tranche metrics
    2. Collateral Pool   — month-by-month pool cash flows
    3. Class A Waterfall — senior tranche distributions
    4. Class B Waterfall — subordinated tranche distributions
    """
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # ---- Sheet 1: Summary ----
        _write_summary_sheet(writer, summary, assumptions_text)

        # ---- Sheet 2: Pool ----
        pool_cf.to_excel(writer, sheet_name="Collateral Pool", index=False)
        _style_sheet(writer.sheets["Collateral Pool"], pool_cf, _POOL_UZS)

        # ---- Sheet 3: Class A ----
        class_a_cf.to_excel(writer, sheet_name="Class A Waterfall", index=False)
        _style_sheet(writer.sheets["Class A Waterfall"], class_a_cf, _TRANCHE_UZS)

        # ---- Sheet 4: Class B ----
        class_b_cf.to_excel(writer, sheet_name="Class B Waterfall", index=False)
        _style_sheet(writer.sheets["Class B Waterfall"], class_b_cf, _TRANCHE_UZS)

    return out_path


def _write_summary_sheet(writer, summary: list[dict], assumptions_text: list[str]):
    """Create the Summary sheet with deal overview and tranche metrics."""
    ws = writer.book.create_sheet("Summary", 0)

    # Title
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "MBS Waterfall Summary / MBS oqim modeli xulosasi"
    title_cell.font = Font(bold=True, size=14, color="1F4E79")

    # Assumptions block
    row = 3
    ws.cell(row=row, column=1, value="Assumptions / Taxminlar").font = _BOLD
    row += 1
    for line in assumptions_text:
        ws.cell(row=row, column=1, value=line)
        row += 1

    # Tranche metrics table
    row += 1
    ws.cell(row=row, column=1, value="Tranche Metrics / Transh ko'rsatkichlari").font = _BOLD
    row += 1

    headers = [
        "Tranche", "Initial Balance (UZS)", "WAL (years)", "Yield (%)",
        "Total Interest (UZS)", "Total Principal (UZS)", "Total Loss (UZS)",
        "Final Balance (UZS)",
    ]
    keys = [
        "tranche", "initial_balance_uzs", "wal_years", "yield_pct",
        "total_interest_uzs", "total_principal_uzs", "total_loss_uzs",
        "final_balance_uzs",
    ]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for s in summary:
        row += 1
        for j, k in enumerate(keys, start=1):
            cell = ws.cell(row=row, column=j, value=s[k])
            if "uzs" in k:
                cell.number_format = _UZS_FMT

    # Column widths
    widths = [14, 22, 12, 12, 22, 22, 18, 18]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
