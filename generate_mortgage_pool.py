"""Generate a sample mortgage pool Excel file with 20 loans for securitization analysis."""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from datetime import date

# Loan data for 20 Uzbekistan mortgage loans
loans = [
    ("LN-001", "Alisher Karimov",    850_000_000, 18.5, 240, 65.0, date(2024, 1, 15), "Tashkent"),
    ("LN-002", "Dilnoza Rashidova",  420_000_000, 19.0, 180, 72.0, date(2024, 2, 10), "Samarkand"),
    ("LN-003", "Bobur Umarov",      1_200_000_000, 17.5, 300, 58.0, date(2023, 11, 5), "Tashkent"),
    ("LN-004", "Gulnara Abdullaeva",  550_000_000, 19.5, 240, 78.0, date(2024, 3, 22), "Bukhara"),
    ("LN-005", "Sardor Mirzaev",      680_000_000, 18.0, 180, 60.0, date(2024, 4, 8),  "Tashkent"),
    ("LN-006", "Malika Nazarova",     350_000_000, 20.0, 120, 80.0, date(2024, 5, 14), "Namangan"),
    ("LN-007", "Jamshid Tursunov",    920_000_000, 17.0, 300, 55.0, date(2023, 9, 30), "Tashkent"),
    ("LN-008", "Feruza Khamidova",    475_000_000, 19.0, 240, 70.0, date(2024, 1, 28), "Fergana"),
    ("LN-009", "Nodir Ismoilov",      780_000_000, 18.5, 180, 63.0, date(2024, 6, 3),  "Tashkent"),
    ("LN-010", "Shahlo Yusupova",     310_000_000, 20.5, 120, 82.0, date(2024, 7, 19), "Andijan"),
    ("LN-011", "Ulugbek Rakhimov",  1_050_000_000, 17.5, 300, 56.0, date(2023, 12, 11),"Tashkent"),
    ("LN-012", "Dilorom Saidova",     620_000_000, 18.0, 240, 68.0, date(2024, 2, 25), "Samarkand"),
    ("LN-013", "Akbar Khodjaev",      890_000_000, 18.5, 180, 61.0, date(2024, 3, 7),  "Tashkent"),
    ("LN-014", "Nigora Tulaganova",   440_000_000, 19.5, 240, 75.0, date(2024, 8, 12), "Nukus"),
    ("LN-015", "Rustam Azimov",       760_000_000, 17.0, 300, 57.0, date(2023, 10, 20),"Tashkent"),
    ("LN-016", "Zulfiya Mamatova",    520_000_000, 19.0, 180, 71.0, date(2024, 4, 30), "Karshi"),
    ("LN-017", "Doniyor Ergashev",    950_000_000, 18.0, 240, 64.0, date(2024, 5, 18), "Tashkent"),
    ("LN-018", "Kamola Ibragimova",   380_000_000, 20.0, 120, 79.0, date(2024, 9, 1),  "Jizzakh"),
    ("LN-019", "Sherzod Ruziev",    1_100_000_000, 17.5, 300, 53.0, date(2024, 1, 5),  "Tashkent"),
    ("LN-020", "Madina Akhmedova",    590_000_000, 19.0, 240, 74.0, date(2024, 6, 22), "Chirchiq"),
]

HEADERS = [
    "Loan ID",
    "Borrower Name",
    "Loan Amount (UZS)",
    "Interest Rate (%)",
    "Term (months)",
    "LTV Ratio (%)",
    "Origination Date",
    "Property City",
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Mortgage Pool"

# ── Styles ──────────────────────────────────────────────────────────────────
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# ── Title row ───────────────────────────────────────────────────────────────
ws.merge_cells("A1:H1")
title_cell = ws["A1"]
title_cell.value = "Sample Mortgage Pool — 20 Loans"
title_cell.font = Font(bold=True, size=14, color="2F5496")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# ── Headers (row 2) ────────────────────────────────────────────────────────
for col_idx, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=2, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border
ws.row_dimensions[2].height = 25

# ── Data rows (rows 3–22) ──────────────────────────────────────────────────
for row_idx, loan in enumerate(loans, start=3):
    for col_idx, value in enumerate(loan, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = thin_border

        # Loan Amount — number format with thousands separator
        if col_idx == 3:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right")
        # Interest Rate / LTV — one decimal
        elif col_idx in (4, 6):
            cell.number_format = '0.0'
            cell.alignment = Alignment(horizontal="center")
        # Term
        elif col_idx == 5:
            cell.alignment = Alignment(horizontal="center")
        # Origination Date
        elif col_idx == 7:
            cell.number_format = 'YYYY-MM-DD'
            cell.alignment = Alignment(horizontal="center")
        # Loan ID / City
        elif col_idx in (1, 8):
            cell.alignment = Alignment(horizontal="center")

# ── Summary section (row 24+) ──────────────────────────────────────────────
summary_start = len(loans) + 4  # blank row then summary

ws.cell(row=summary_start, column=1, value="Pool Summary").font = Font(bold=True, size=12, color="2F5496")

labels_values = [
    ("Number of Loans",    len(loans)),
    ("Total Pool Balance", sum(l[2] for l in loans)),
    ("Average Loan Size",  sum(l[2] for l in loans) / len(loans)),
    ("Avg Interest Rate",  sum(l[3] for l in loans) / len(loans)),
    ("Avg LTV Ratio",      sum(l[5] for l in loans) / len(loans)),
    ("Avg Term (months)",  sum(l[4] for l in loans) / len(loans)),
]

for i, (label, val) in enumerate(labels_values, start=1):
    row = summary_start + i
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = Font(bold=True)
    val_cell = ws.cell(row=row, column=2, value=val)
    if "Balance" in label or "Size" in label:
        val_cell.number_format = '#,##0'
    elif "Rate" in label or "LTV" in label:
        val_cell.number_format = '0.0'
    elif "Term" in label:
        val_cell.number_format = '0.0'

# ── Column widths ───────────────────────────────────────────────────────────
col_widths = [12, 24, 22, 16, 14, 14, 18, 16]
for idx, width in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = width

# ── Freeze panes below header ──────────────────────────────────────────────
ws.freeze_panes = "A3"

# ── Save ────────────────────────────────────────────────────────────────────
output_path = "mortgage_pool.xlsx"
wb.save(output_path)
print(f"Created {output_path} with {len(loans)} loans.")
