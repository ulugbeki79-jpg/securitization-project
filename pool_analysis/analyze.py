"""Mortgage loan pool analysis for MBS structuring.

Reads an Excel loan pool, computes weighted-average metrics, generates
distribution charts, flags ineligible loans, and outputs a Word report.

Usage:
    python -m pool_analysis.analyze sample_data/loan_pool.xlsx

The Word report and chart images are written to the output/ directory.
"""

from __future__ import annotations

import argparse
import os
import textwrap
from dataclasses import dataclass
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # non-interactive backend — must be set before pyplot import
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

# ---------------------------------------------------------------------------
# Configuration — eligibility criteria
# ---------------------------------------------------------------------------

MAX_LTV_PCT = 80.0
MAX_DAYS_PAST_DUE = 30


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

REQUIRED_COLUMNS = {
    "loan_id",
    "region",
    "principal_uzs",
    "property_value_uzs",
    "ltv_pct",
    "coupon_pct",
    "original_term_months",
    "remaining_term_months",
    "days_past_due",
}


def load_pool(path: str | Path) -> pd.DataFrame:
    """Read the first sheet of an Excel workbook and validate columns."""
    df = pd.read_excel(path, sheet_name=0)
    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    return df


# ---------------------------------------------------------------------------
# Weighted-average metrics
# ---------------------------------------------------------------------------

@dataclass
class PoolMetrics:
    loan_count: int
    total_principal_uzs: float
    wac: float           # weighted average coupon (%)
    wam: float           # weighted average maturity (months)
    waltv: float         # weighted average LTV (%)
    avg_loan_size_uzs: float
    min_loan_uzs: float
    max_loan_uzs: float


def compute_metrics(df: pd.DataFrame) -> PoolMetrics:
    total = df["principal_uzs"].sum()
    weights = df["principal_uzs"] / total
    return PoolMetrics(
        loan_count=len(df),
        total_principal_uzs=total,
        wac=round((weights * df["coupon_pct"]).sum(), 4),
        wam=round((weights * df["remaining_term_months"]).sum(), 2),
        waltv=round((weights * df["ltv_pct"]).sum(), 2),
        avg_loan_size_uzs=round(total / len(df)),
        min_loan_uzs=df["principal_uzs"].min(),
        max_loan_uzs=df["principal_uzs"].max(),
    )


# ---------------------------------------------------------------------------
# Eligibility screening
# ---------------------------------------------------------------------------

def flag_ineligible(df: pd.DataFrame) -> pd.DataFrame:
    """Return a copy of the pool with an ``eligible`` column and reason flags."""
    out = df.copy()
    out["flag_ltv"] = out["ltv_pct"] > MAX_LTV_PCT
    out["flag_past_due"] = out["days_past_due"] > MAX_DAYS_PAST_DUE
    out["eligible"] = ~(out["flag_ltv"] | out["flag_past_due"])

    reasons = []
    for _, row in out.iterrows():
        r = []
        if row["flag_ltv"]:
            r.append(f"LTV {row['ltv_pct']:.1f}% > {MAX_LTV_PCT}%")
        if row["flag_past_due"]:
            r.append(f"Past due {int(row['days_past_due'])}d > {MAX_DAYS_PAST_DUE}d")
        reasons.append("; ".join(r) if r else "")
    out["ineligibility_reason"] = reasons
    return out


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------

def _format_uzs(value: float, _pos=None) -> str:
    """Format large UZS values as human-readable (e.g. 150 mln)."""
    if abs(value) >= 1e9:
        return f"{value / 1e9:.1f} mlrd"
    if abs(value) >= 1e6:
        return f"{value / 1e6:.0f} mln"
    return f"{value:,.0f}"


def chart_loan_size_distribution(df: pd.DataFrame, out_dir: Path) -> Path:
    bins = [0, 100e6, 200e6, 300e6, 500e6, 800e6, float("inf")]
    labels = ["<100 mln", "100-200 mln", "200-300 mln", "300-500 mln", "500-800 mln", "800+ mln"]
    df = df.copy()
    df["size_bucket"] = pd.cut(df["principal_uzs"], bins=bins, labels=labels, right=True)
    counts = df["size_bucket"].value_counts().reindex(labels, fill_value=0)

    fig, ax = plt.subplots(figsize=(8, 4.5))
    bars = ax.bar(counts.index, counts.values, color="#2E86AB", edgecolor="white")
    ax.set_title("Loan Size Distribution (UZS)", fontsize=13, fontweight="bold")
    ax.set_ylabel("Number of Loans")
    ax.set_xlabel("Principal Bucket")
    for bar, val in zip(bars, counts.values):
        if val > 0:
            ax.text(bar.get_x() + bar.get_width() / 2, val + 0.5, str(val),
                    ha="center", va="bottom", fontsize=9)
    plt.xticks(rotation=30, ha="right")
    plt.tight_layout()
    path = out_dir / "chart_loan_size.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def chart_region_distribution(df: pd.DataFrame, out_dir: Path) -> Path:
    counts = df["region"].value_counts().sort_values(ascending=True)

    fig, ax = plt.subplots(figsize=(8, max(4.5, len(counts) * 0.35)))
    ax.barh(counts.index, counts.values, color="#A23B72", edgecolor="white")
    ax.set_title("Loans by Region", fontsize=13, fontweight="bold")
    ax.set_xlabel("Number of Loans")
    for i, val in enumerate(counts.values):
        ax.text(val + 0.3, i, str(val), va="center", fontsize=9)
    plt.tight_layout()
    path = out_dir / "chart_region.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def chart_ltv_distribution(df: pd.DataFrame, out_dir: Path) -> Path:
    bins = [0, 50, 60, 70, 80, 90, 100]
    labels = ["≤50%", "50-60%", "60-70%", "70-80%", "80-90%", "90-100%"]
    df = df.copy()
    df["ltv_bucket"] = pd.cut(df["ltv_pct"], bins=bins, labels=labels, right=True)
    counts = df["ltv_bucket"].value_counts().reindex(labels, fill_value=0)

    colors = ["#2ca02c", "#2ca02c", "#FFB627", "#FFB627", "#E63946", "#E63946"]
    fig, ax = plt.subplots(figsize=(7, 4.5))
    bars = ax.bar(counts.index, counts.values, color=colors, edgecolor="white")
    ax.set_title("LTV Distribution", fontsize=13, fontweight="bold")
    ax.set_ylabel("Number of Loans")
    ax.set_xlabel("LTV Bucket")
    ax.axvline(x=3.5, color="red", linestyle="--", linewidth=1, label=f"Eligibility cutoff ({MAX_LTV_PCT}%)")
    ax.legend(fontsize=9)
    for bar, val in zip(bars, counts.values):
        if val > 0:
            ax.text(bar.get_x() + bar.get_width() / 2, val + 0.5, str(val),
                    ha="center", va="bottom", fontsize=9)
    plt.tight_layout()
    path = out_dir / "chart_ltv.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


# ---------------------------------------------------------------------------
# Word report
# ---------------------------------------------------------------------------

def _add_heading(doc: Document, text_en: str, text_uz: str, level: int = 1):
    doc.add_heading(text_en, level=level)
    p = doc.add_paragraph()
    run = p.add_run(text_uz)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = None  # inherit


def _add_metric_table(doc: Document, metrics: PoolMetrics):
    rows = [
        ("Total Loans / Jami kreditlar", f"{metrics.loan_count:,}"),
        ("Total Principal (UZS) / Jami asosiy qarz", f"{metrics.total_principal_uzs:,.0f}"),
        ("Average Loan Size (UZS) / O'rtacha kredit", f"{metrics.avg_loan_size_uzs:,.0f}"),
        ("Min Loan (UZS) / Eng kichik kredit", f"{metrics.min_loan_uzs:,.0f}"),
        ("Max Loan (UZS) / Eng katta kredit", f"{metrics.max_loan_uzs:,.0f}"),
        ("WAC (%) / O'rtacha vaznli kupon", f"{metrics.wac:.2f}%"),
        ("WAM (months) / O'rtacha vaznli muddat", f"{metrics.wam:.1f}"),
        ("WALTV (%) / O'rtacha vaznli LTV", f"{metrics.waltv:.2f}%"),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Metric / Ko'rsatkich"
    hdr[1].text = "Value / Qiymat"
    for i, (label, value) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value


def _add_ineligible_table(doc: Document, flagged: pd.DataFrame):
    bad = flagged[~flagged["eligible"]].sort_values("loan_id")
    if bad.empty:
        doc.add_paragraph("All loans meet eligibility criteria. / Barcha kreditlar mezonlarga mos.")
        return

    doc.add_paragraph(
        f"{len(bad)} loan(s) flagged as ineligible / {len(bad)} ta kredit mezonlarga mos emas:"
    )
    table = doc.add_table(rows=len(bad) + 1, cols=5)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Loan ID", "Region", "LTV %", "Days Past Due", "Reason / Sabab"]
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, (_, row) in enumerate(bad.iterrows(), start=1):
        table.rows[i].cells[0].text = str(row["loan_id"])
        table.rows[i].cells[1].text = str(row["region"])
        table.rows[i].cells[2].text = f"{row['ltv_pct']:.1f}"
        table.rows[i].cells[3].text = str(int(row["days_past_due"]))
        table.rows[i].cells[4].text = str(row["ineligibility_reason"])


def generate_report(
    metrics: PoolMetrics,
    flagged: pd.DataFrame,
    chart_paths: dict[str, Path],
    out_dir: Path,
) -> Path:
    doc = Document()

    # Title
    doc.add_heading("Mortgage Loan Pool Analysis", level=0)
    p = doc.add_paragraph()
    run = p.add_run("Ipoteka kredit hovuzi tahlili")
    run.italic = True
    run.font.size = Pt(11)

    # --- Pool summary ---
    _add_heading(doc, "Pool Summary", "Hovuz xulosasi", level=1)
    _add_metric_table(doc, metrics)

    # --- Eligibility criteria ---
    _add_heading(doc, "Eligibility Criteria", "Muvofiqlik mezonlari", level=1)
    doc.add_paragraph(
        f"• Maximum LTV: {MAX_LTV_PCT}%\n"
        f"• Maximum Days Past Due: {MAX_DAYS_PAST_DUE}"
    )

    eligible_count = flagged["eligible"].sum()
    total = len(flagged)
    doc.add_paragraph(
        f"Eligible: {eligible_count} / {total} loans "
        f"({eligible_count / total * 100:.1f}%)"
    )

    # --- Charts ---
    _add_heading(doc, "Distribution Analysis", "Taqsimot tahlili", level=1)

    for title, key in [
        ("Loan Size Distribution / Kredit hajmi taqsimoti", "loan_size"),
        ("Regional Distribution / Hududiy taqsimot", "region"),
        ("LTV Distribution / LTV taqsimoti", "ltv"),
    ]:
        doc.add_heading(title, level=2)
        doc.add_picture(str(chart_paths[key]), width=Inches(5.8))

    # --- Ineligible loans ---
    _add_heading(doc, "Ineligible Loans", "Nomuvofiq kreditlar", level=1)
    _add_ineligible_table(doc, flagged)

    path = out_dir / "loan_pool_report.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def run(excel_path: str, out_dir: str = "output") -> None:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    print(f"Reading pool from {excel_path} ...")
    df = load_pool(excel_path)
    print(f"  Loaded {len(df)} loans")

    print("Computing weighted-average metrics ...")
    metrics = compute_metrics(df)
    print(f"  WAC  = {metrics.wac:.2f}%")
    print(f"  WAM  = {metrics.wam:.1f} months")
    print(f"  WALTV = {metrics.waltv:.2f}%")

    print("Screening eligibility ...")
    flagged = flag_ineligible(df)
    n_bad = (~flagged["eligible"]).sum()
    print(f"  {n_bad} loan(s) flagged ineligible out of {len(df)}")

    print("Generating charts ...")
    charts = {
        "loan_size": chart_loan_size_distribution(df, out),
        "region": chart_region_distribution(df, out),
        "ltv": chart_ltv_distribution(df, out),
    }

    print("Writing Word report ...")
    report_path = generate_report(metrics, flagged, charts, out)
    print(f"  Report saved to {report_path}")

    # Also save the flagged pool back to Excel for reference
    flagged_path = out / "loan_pool_flagged.xlsx"
    flagged.to_excel(flagged_path, index=False, sheet_name="Flagged Pool")
    print(f"  Flagged pool saved to {flagged_path}")

    print("Done.")


def main():
    parser = argparse.ArgumentParser(
        description="Analyze a mortgage loan pool for MBS structuring",
    )
    parser.add_argument(
        "excel_file",
        help="Path to the Excel file containing the loan pool",
    )
    parser.add_argument(
        "-o", "--output-dir",
        default="output",
        help="Directory for report and charts (default: output/)",
    )
    args = parser.parse_args()
    run(args.excel_file, args.output_dir)


if __name__ == "__main__":
    main()
